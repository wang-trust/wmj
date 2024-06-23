#coding=utf-8


filepath = 'C:\\Users\\goupi\\Desktop\\electron\\demo11\\dist\\'


# read docx
import docx

def readdocx():
    doc = docx.Document(filepath + '12.docx')
    for each in doc.paragraphs:
        print(each.text)


# read doc
from win32com import client
import os

# 转换doc为docx
def doc2docx():
    fn = filepath + '12.doc'
    word = client.Dispatch("Word.Application")  # 打开word应用程序
    doc = word.Documents.Open(fn)  # 打开word文件

    a = os.path.split(fn)  # 分离路径和文件
    b = os.path.splitext(a[-1])[0]  # 拿到文件名

    doc.SaveAs("{}\\{}.docx".format(a[0], b), 12)  # 另存为后缀为".docx"的文件，其中参数12或16指docx文件
    doc.Close()  # 关闭原来word文件
    word.Quit()

# read .xlsx
import openpyxl

def readxlsx():
    workbook = openpyxl.load_workbook(filepath + '123.xlsx')

# 选择工作表
    sheet = workbook['Sheet1']
    cell_value = sheet['A1'].value
    print(cell_value)


# read .xls
import xlrd
# 打开文件
def readxls():
    data = xlrd.open_workbook(filepath + '12.xls')
    table = data.sheet_by_index(0)
    v = table.cell(0, 0).value
    print(v)



if __name__ == '__main__':
    readxls()


